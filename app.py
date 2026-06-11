from flask import Flask, render_template, request, jsonify, send_file, redirect, url_for, session
from werkzeug.utils import secure_filename
from docx import Document
from num2words import num2words
from openpyxl import Workbook, load_workbook
import sqlite3
import os
import io
import shutil
from datetime import datetime, date, timedelta

app = Flask(__name__)
app.secret_key = 'exceledge_hrms_secret_key_2026'  # change in production
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_FILE = os.path.join(BASE_DIR, "hrms.db")
TEMPLATE_PATH = os.path.join(BASE_DIR, "_Offer_Letter_BSR.docx")
LEAVE_TYPES = ["Earned Leave", "Casual Leave", "Loss of Pay", "Work From Home"]
UPLOAD_FOLDER = os.path.join(BASE_DIR, "uploads")
os.makedirs(UPLOAD_FOLDER, exist_ok=True)


def get_connection():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    return conn


def get_next_employee_code():
    """Generate next employee code based on the last one (e.g. EE/69 -> EE/70)"""
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT employee_code FROM employees")
    codes = cursor.fetchall()
    conn.close()
    
    max_num = 69  # base on last given EE/69
    for row in codes:
        code = row[0] if row[0] else ''
        if code and code.startswith('EE/'):
            try:
                num = int(code.split('/')[-1])
                if num > max_num:
                    max_num = num
            except:
                pass
    return f"EE/{max_num + 1}"


def ensure_column(cursor, table, column, definition):
    cursor.execute(f"PRAGMA table_info({table})")
    existing = [row["name"] for row in cursor.fetchall()]
    if column not in existing:
        cursor.execute(f"ALTER TABLE {table} ADD COLUMN {column} {definition}")


def init_db():
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS departments ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "name TEXT UNIQUE NOT NULL,"
        "description TEXT"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS offices ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "name TEXT UNIQUE NOT NULL,"
        "location TEXT"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS designations ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "title TEXT NOT NULL,"
        "department_id INTEGER"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS shifts ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "name TEXT NOT NULL,"
        "start_time TEXT,"
        "end_time TEXT,"
        "description TEXT"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS employees ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_code TEXT UNIQUE NOT NULL,"
        "first_name TEXT NOT NULL,"
        "last_name TEXT NOT NULL,"
        "email TEXT,"
        "phone TEXT,"
        "dob TEXT,"
        "gender TEXT,"
        "department_id INTEGER,"
        "designation_id INTEGER,"
        "office_id INTEGER,"
        "manager_id INTEGER,"
        "joining_date TEXT,"
        "ctc REAL,"
        "status TEXT,"
        "address TEXT,"
        "pf_active INTEGER DEFAULT 1,"
        "esic_active INTEGER DEFAULT 1,"
        "professional_tax_active INTEGER DEFAULT 1,"
        "lwf_active INTEGER DEFAULT 1"
        ")"
    )
    ensure_column(cursor, "employees", "pf_active", "INTEGER DEFAULT 1")
    ensure_column(cursor, "employees", "esic_active", "INTEGER DEFAULT 1")
    ensure_column(cursor, "employees", "professional_tax_active", "INTEGER DEFAULT 1")
    ensure_column(cursor, "employees", "lwf_active", "INTEGER DEFAULT 1")
    ensure_column(cursor, "employees", "password", "TEXT")

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS employee_joining_forms ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "full_name TEXT,"
        "contact_no TEXT,"
        "email_id TEXT,"
        "designation TEXT,"
        "joining_date TEXT,"
        "permanent_address TEXT,"
        "date_of_birth TEXT,"
        "gender TEXT,"
        "marital_status TEXT,"
        "pan_no TEXT,"
        "aadhar_no TEXT,"
        "bank_name TEXT,"
        "bank_account_number TEXT,"
        "ifsc_code TEXT,"
        "branch_name TEXT,"
        "highest_qualification TEXT,"
        "university_board TEXT,"
        "year_of_passing TEXT,"
        "percentage_grade TEXT,"
        "emergency_person_name TEXT,"
        "emergency_person_mobile TEXT,"
        "emergency_relation TEXT,"
        "emergency_person_address TEXT,"
        "prev_company_name TEXT,"
        "prev_designation TEXT,"
        "prev_duration TEXT,"
        "prev_last_salary TEXT,"
        "photo_filename TEXT,"
        "pan_card_filename TEXT,"
        "aadhar_card_filename TEXT,"
        "cheque_passbook_filename TEXT,"
        "highest_education_cert_filename TEXT,"
        "last_3_month_salary_slip_filename TEXT,"
        "prev_employment_docs_filename TEXT,"
        "submitted_on TEXT"
        ")"
    )

    # For existing DBs with old schema, add new columns
    new_joining_cols = [
        ("highest_qualification", "TEXT"),
        ("university_board", "TEXT"),
        ("year_of_passing", "TEXT"),
        ("percentage_grade", "TEXT"),
        ("emergency_person_name", "TEXT"),
        ("emergency_person_mobile", "TEXT"),
        ("emergency_relation", "TEXT"),
        ("emergency_person_address", "TEXT"),
        ("prev_company_name", "TEXT"),
        ("prev_designation", "TEXT"),
        ("prev_duration", "TEXT"),
        ("prev_last_salary", "TEXT"),
        ("photo_filename", "TEXT"),
        ("pan_card_filename", "TEXT"),
        ("aadhar_card_filename", "TEXT"),
        ("cheque_passbook_filename", "TEXT"),
        ("highest_education_cert_filename", "TEXT"),
        ("last_3_month_salary_slip_filename", "TEXT"),
        ("prev_employment_docs_filename", "TEXT"),
    ]
    for col, defn in new_joining_cols:
        ensure_column(cursor, "employee_joining_forms", col, defn)

    ensure_column(cursor, "employee_joining_forms", "employee_code", "TEXT")
    ensure_column(cursor, "employee_joining_forms", "status", "TEXT DEFAULT 'pending'")

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS attendance ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_id INTEGER NOT NULL,"
        "date TEXT NOT NULL,"
        "status TEXT NOT NULL,"
        "shift_id INTEGER,"
        "check_in TEXT,"
        "check_out TEXT"
        ")"
    )
    ensure_column(cursor, "attendance", "remarks", "TEXT")

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS leave_balances ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_id INTEGER NOT NULL,"
        "earned_leave REAL DEFAULT 0,"
        "casual_leave REAL DEFAULT 0,"
        "loss_of_pay REAL DEFAULT 0,"
        "year INTEGER NOT NULL,"
        "FOREIGN KEY (employee_id) REFERENCES employees (id)"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS leave_requests ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_id INTEGER NOT NULL,"
        "leave_type TEXT NOT NULL,"
        "start_date TEXT NOT NULL,"
        "end_date TEXT NOT NULL,"
        "status TEXT DEFAULT 'Pending',"
        "reason TEXT,"
        "applied_on TEXT,"
        "FOREIGN KEY (employee_id) REFERENCES employees (id)"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS earned_leave_credits ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_id INTEGER NOT NULL,"
        "year INTEGER NOT NULL,"
        "month INTEGER NOT NULL,"
        "credited REAL NOT NULL,"
        "credited_on TEXT,"
        "UNIQUE(employee_id, year, month),"
        "FOREIGN KEY (employee_id) REFERENCES employees (id)"
        ")"
    )

    cursor.execute(
        "CREATE TABLE IF NOT EXISTS payrolls ("
        "id INTEGER PRIMARY KEY AUTOINCREMENT,"
        "employee_id INTEGER NOT NULL,"
        "month INTEGER NOT NULL,"
        "year INTEGER NOT NULL,"
        "ctc REAL NOT NULL,"
        "basic REAL NOT NULL,"
        "da REAL NOT NULL,"
        "hra REAL NOT NULL,"
        "conveyance REAL NOT NULL,"
        "special_allowance REAL NOT NULL,"
        "gross REAL NOT NULL,"
        "pf REAL NOT NULL,"
        "esic REAL NOT NULL,"
        "professional_tax REAL NOT NULL,"
        "lwf REAL NOT NULL,"
        "total_deductions REAL NOT NULL,"
        "net_pay REAL NOT NULL"
        ")"
    )

    default_departments = [
        ("Human Resources", "People operations and employee success."),
        ("Engineering", "Product and software development."),
        ("Finance", "Payroll, accounts and budgeting."),
        ("Sales", "Revenue generation and customer success.")
    ]
    for name, description in default_departments:
        try:
            cursor.execute("INSERT INTO departments (name, description) VALUES (?, ?)", (name, description))
        except sqlite3.IntegrityError:
            pass

    default_offices = [
        ("Head Office", "Downtown Campus"),
        ("Regional Office", "Sector 21")
    ]
    for name, location in default_offices:
        try:
            cursor.execute("INSERT INTO offices (name, location) VALUES (?, ?)", (name, location))
        except sqlite3.IntegrityError:
            pass

    default_designations = [
        ("HR Manager", 1),
        ("Software Engineer", 2),
        ("Accountant", 3),
        ("Sales Executive", 4)
    ]
    for title, dept_id in default_designations:
        cursor.execute(
            "SELECT id FROM designations WHERE title = ? AND department_id = ?",
            (title, dept_id)
        )
        if not cursor.fetchone():
            cursor.execute(
                "INSERT INTO designations (title, department_id) VALUES (?, ?)",
                (title, dept_id)
            )

    default_shifts = [
        ("Day Shift", "09:00", "18:00", "Standard daytime shift."),
        ("Night Shift", "22:00", "06:00", "Overnight coverage shift.")
    ]
    for name, start_time, end_time, description in default_shifts:
        cursor.execute(
            "SELECT id FROM shifts WHERE name = ?",
            (name,)
        )
        if not cursor.fetchone():
            cursor.execute(
                "INSERT INTO shifts (name, start_time, end_time, description) VALUES (?, ?, ?, ?)",
                (name, start_time, end_time, description)
            )

    conn.commit()
    conn.close()


def ensure_template():
    if os.path.exists(TEMPLATE_PATH):
        return

    doc = Document()
    doc.add_paragraph("Date: [DATE]")
    doc.add_paragraph("")
    doc.add_paragraph("Dear [PREFIX] [NAME],")
    doc.add_paragraph("")
    doc.add_paragraph("We are pleased to offer you a position as [POSITION] at our company.")
    doc.add_paragraph("")
    doc.add_paragraph("Position Details:")
    doc.add_paragraph("Joining Date: [JOINING_DATE]")
    doc.add_paragraph("Location: [LOCATION]")
    doc.add_paragraph("Department: [DEPARTMENT]")
    doc.add_paragraph("Monthly CTC: [CTC_M] ([CTC_M_WORD])")
    doc.add_paragraph("Annual CTC: [CTC_A] ([CTC_A_WORD])")
    doc.add_paragraph("")
    doc.add_paragraph("We look forward to welcoming you to the team.")
    doc.add_paragraph("")
    doc.add_paragraph("Best regards,")
    doc.add_paragraph("HR Team")
    doc.save(TEMPLATE_PATH)


def replace_placeholders(doc, replacements, bold_keys=None):
    if bold_keys is None:
        bold_keys = []

    def process_paragraph(paragraph):
        full_text = "".join(run.text for run in paragraph.runs)
        if any(key in full_text for key in replacements):
            for key, value in replacements.items():
                full_text = full_text.replace(key, value)
            for run in list(paragraph.runs):
                paragraph._element.remove(run._element)

            bold_ranges = []
            for bold_key in bold_keys:
                value = replacements.get(bold_key, "")
                if value and value in full_text:
                    start = full_text.find(value)
                    if start != -1:
                        bold_ranges.append((start, start + len(value)))
            bold_ranges.sort()

            cursor = 0
            for start, end in bold_ranges:
                if start > cursor:
                    paragraph.add_run(full_text[cursor:start])
                paragraph.add_run(full_text[start:end]).bold = True
                cursor = end
            if cursor < len(full_text):
                paragraph.add_run(full_text[cursor:])

    for paragraph in doc.paragraphs:
        process_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    process_paragraph(paragraph)
    for section in doc.sections:
        for header in section.header.paragraphs:
            process_paragraph(header)
        for footer in section.footer.paragraphs:
            process_paragraph(footer)


def dict_rows(rows):
    return [dict(row) for row in rows]


def apply_weekend_absence_rule(records, start_date=None, end_date=None):
    if not records:
        return records

    def normalize_date(value):
        try:
            return datetime.strptime(value, '%Y-%m-%d').date()
        except Exception:
            return None

    def weekday_score(status):
        if status in ('P', 'H', 'WO'):
            return 1.0
        if status in ('AH', 'HF'):
            return 0.5
        return 0.0

    if start_date:
        try:
            start_dt = datetime.strptime(start_date, '%Y-%m-%d').date()
        except Exception:
            start_dt = None
    else:
        start_dt = None
    if end_date:
        try:
            end_dt = datetime.strptime(end_date, '%Y-%m-%d').date()
        except Exception:
            end_dt = None
    else:
        end_dt = None

    by_employee = {}
    for row in records:
        emp_id = row.get('employee_id')
        if emp_id is None:
            continue
        dt = normalize_date(row.get('date'))
        if not dt:
            continue
        week_key = (dt.isocalendar()[0], dt.isocalendar()[1])
        by_employee.setdefault(emp_id, {})
        by_employee[emp_id].setdefault(week_key, {'records': {}})
        by_employee[emp_id][week_key]['records'][dt] = row

    extra_rows = []
    for emp_id, weeks in by_employee.items():
        for week_key, week_data in weeks.items():
            year, week_num = week_key
            try:
                monday = datetime.fromisocalendar(year, week_num, 1).date()
            except ValueError:
                continue
            total_score = 0.0
            for offset in range(5):
                day = monday + timedelta(days=offset)
                row = week_data['records'].get(day)
                status = row.get('status') if row else None
                total_score += weekday_score(status)
            if total_score >= 3.5:
                continue
            for offset in (5, 6):
                day = monday + timedelta(days=offset)
                if start_dt and day < start_dt:
                    continue
                if end_dt and day > end_dt:
                    continue
                existing = week_data['records'].get(day)
                if existing:
                    if existing.get('status') != 'A':
                        existing['status'] = 'A'
                    continue
                extra_rows.append({
                    'id': None,
                    'employee_id': emp_id,
                    'date': day.strftime('%Y-%m-%d'),
                    'status': 'A',
                    'shift_name': None,
                    'check_in': None,
                    'check_out': None
                })

    # Additional conditions for WO and H to A
    for emp_id, weeks in by_employee.items():
        week_keys = sorted(weeks.keys())
        for i, week_key in enumerate(week_keys):
            year, week_num = week_key
            monday = datetime.fromisocalendar(year, week_num, 1).date()
            friday = monday + timedelta(days=4)
            sat = monday + timedelta(days=5)
            sun = monday + timedelta(days=6)
            friday_row = weeks[week_key]['records'].get(friday)
            friday_status = friday_row.get('status') if friday_row else None
            next_monday = monday + timedelta(days=7)
            next_week_key = (next_monday.isocalendar()[0], next_monday.isocalendar()[1])
            next_monday_row = weeks.get(next_week_key, {}).get('records', {}).get(next_monday)
            next_monday_status = next_monday_row.get('status') if next_monday_row else None

            # Condition 1: absent on Friday and Monday -> WO to A
            if friday_status == 'A' and next_monday_status == 'A':
                for day in [sat, sun]:
                    row = weeks[week_key]['records'].get(day)
                    if row and row.get('status') == 'WO':
                        row['status'] = 'A'

            # Condition 2: holiday on Friday with absences around -> WO and H to A
            if friday_status == 'H':
                thursday = monday + timedelta(days=3)
                thursday_row = weeks[week_key]['records'].get(thursday)
                thursday_status = thursday_row.get('status') if thursday_row else None
                tuesday = next_monday + timedelta(days=1)
                tuesday_row = weeks.get(next_week_key, {}).get('records', {}).get(tuesday)
                tuesday_status = tuesday_row.get('status') if tuesday_row else None
                absent_around = (thursday_status == 'A' or friday_status == 'A') and (next_monday_status == 'A' or tuesday_status == 'A')
                if absent_around:
                    for day in [sat, sun]:
                        row = weeks[week_key]['records'].get(day)
                        if row and row.get('status') == 'WO':
                            row['status'] = 'A'
                    if friday_row:
                        friday_row['status'] = 'A'

            # Condition 3: holiday on Monday with absences around -> WO and H to A
            if next_monday_status == 'H':
                saturday = monday + timedelta(days=5)
                saturday_row = weeks[week_key]['records'].get(saturday)
                saturday_status = saturday_row.get('status') if saturday_row else None
                tuesday = next_monday + timedelta(days=1)
                tuesday_row = weeks.get(next_week_key, {}).get('records', {}).get(tuesday)
                tuesday_status = tuesday_row.get('status') if tuesday_row else None
                wednesday = next_monday + timedelta(days=2)
                wednesday_row = weeks.get(next_week_key, {}).get('records', {}).get(wednesday)
                wednesday_status = wednesday_row.get('status') if wednesday_row else None
                absent_around = (friday_status == 'A' or saturday_status == 'A') and (tuesday_status == 'A' or wednesday_status == 'A')
                if absent_around:
                    for day in [sat, sun]:
                        row = weeks[week_key]['records'].get(day)
                        if row and row.get('status') == 'WO':
                            row['status'] = 'A'
                    if next_monday_row:
                        next_monday_row['status'] = 'A'

    return records + extra_rows


def calculate_pay_components(ctc, pf_active=True, esic_active=True, professional_tax_active=True, lwf_active=True):
    basic = round(ctc * 0.30, 2)
    da = round(ctc * 0.20, 2)
    hra = round(ctc * 0.30, 2)
    conveyance = round(ctc * 0.10, 2)
    special_allowance = round(ctc * 0.10, 2)
    gross = round(basic + da + hra + conveyance + special_allowance, 2)
    pf = round(basic * 0.12, 2) if pf_active else 0.0
    esic = round(gross * 0.0475, 2) if esic_active else 0.0
    professional_tax = 200.0 if professional_tax_active else 0.0
    lwf = 30.0 if lwf_active else 0.0
    total_deductions = round(pf + esic + professional_tax + lwf, 2)
    net_pay = round(gross - total_deductions, 2)
    return {
        "basic": basic,
        "da": da,
        "hra": hra,
        "conveyance": conveyance,
        "special_allowance": special_allowance,
        "gross": gross,
        "pf": pf,
        "esic": esic,
        "professional_tax": professional_tax,
        "lwf": lwf,
        "total_deductions": total_deductions,
        "net_pay": net_pay
    }


def get_employee_list():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT e.*, d.name AS department_name, des.title AS designation_title, o.name AS office_name, m.first_name || ' ' || m.last_name AS manager_name "
        "FROM employees e "
        "LEFT JOIN departments d ON e.department_id = d.id "
        "LEFT JOIN designations des ON e.designation_id = des.id "
        "LEFT JOIN offices o ON e.office_id = o.id "
        "LEFT JOIN employees m ON e.manager_id = m.id "
        "ORDER BY e.first_name, e.last_name"
    )
    rows = cursor.fetchall()
    conn.close()
    return dict_rows(rows)


def get_lookup_data():
    conn = get_connection()
    cursor = conn.cursor()

    cursor.execute("SELECT * FROM departments ORDER BY name")
    departments = dict_rows(cursor.fetchall())
    cursor.execute("SELECT * FROM offices ORDER BY name")
    offices = dict_rows(cursor.fetchall())
    cursor.execute("SELECT * FROM designations ORDER BY title")
    designations = dict_rows(cursor.fetchall())
    cursor.execute("SELECT * FROM shifts ORDER BY name")
    shifts = dict_rows(cursor.fetchall())
    employees = get_employee_list()

    conn.close()
    return {
        "departments": departments,
        "offices": offices,
        "designations": designations,
        "shifts": shifts,
        "employees": employees,
        "leave_types": LEAVE_TYPES
    }

@app.route("/")
def index():
    if not session.get('logged_in'):
        return redirect(url_for('login_page'))
    return render_template("index.html")

@app.route("/login", methods=["GET", "POST"])
def login_page():
    if request.method == "POST":
        user_id = request.form.get("user_id", "").strip()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")

        if user_id == "admin" and password == "admin":
            session['logged_in'] = True
            session['is_admin'] = True
            session['employee_id'] = None
            return jsonify({"success": True})

        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM employees WHERE phone = ?", (user_id,))
        emp = cursor.fetchone()
        conn.close()

        if not emp:
            return jsonify({"success": False, "error": "Employee not found with this mobile number."})

        stored_pass = emp['password'] or ""
        if not stored_pass:
            # First time login - set the password
            if confirm and password != confirm:
                return jsonify({"success": False, "error": "Passwords do not match.", "first_time": True})
            conn = get_connection()
            cursor = conn.cursor()
            cursor.execute("UPDATE employees SET password = ? WHERE id = ?", (password, emp['id']))
            conn.commit()
            conn.close()
            session['logged_in'] = True
            session['is_admin'] = False
            session['employee_id'] = emp['id']
            session['user_phone'] = user_id
            return jsonify({"success": True})

        if stored_pass == password:
            session['logged_in'] = True
            session['is_admin'] = False
            session['employee_id'] = emp['id']
            session['user_phone'] = user_id
            return jsonify({"success": True})
        else:
            return jsonify({"success": False, "error": "Invalid password."})

    # GET - show login page
    if session.get('logged_in'):
        return redirect(url_for('index'))
    return render_template("login.html")

@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for('login_page'))

@app.route("/api/current_user")
def current_user():
    if not session.get('logged_in'):
        return jsonify({"error": "Not logged in"}), 401
    if session.get('is_admin'):
        return jsonify({"is_admin": True, "employee_id": None})
    emp_id = session.get('employee_id')
    if emp_id:
        conn = get_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id, employee_code, first_name, last_name, phone FROM employees WHERE id = ?", (emp_id,))
        emp = cursor.fetchone()
        conn.close()
        if emp:
            return jsonify({
                "is_admin": False,
                "employee_id": emp['id'],
                "phone": emp['phone'],
                "name": f"{emp['first_name']} {emp['last_name']}",
                "employee_code": emp['employee_code']
            })
    return jsonify({"is_admin": False, "employee_id": None})

@app.route("/api/lookups")
def lookups():
    return jsonify(get_lookup_data())

@app.route("/api/departments", methods=["GET", "POST"])
def manage_departments():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM departments ORDER BY name")
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    name = payload.get("name", "").strip()
    description = payload.get("description", "").strip()
    if not name:
        conn.close()
        return jsonify({"error": "Department name is required."}), 400
    try:
        cursor.execute("INSERT INTO departments (name, description) VALUES (?, ?)", (name, description))
        conn.commit()
        return jsonify({"success": True})
    except sqlite3.IntegrityError:
        return jsonify({"error": "Department already exists."}), 400
    finally:
        conn.close()

@app.route("/api/departments/<int:department_id>", methods=["DELETE"])
def delete_department(department_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM departments WHERE id = ?", (department_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/offices", methods=["GET", "POST"])
def manage_offices():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM offices ORDER BY name")
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    name = payload.get("name", "").strip()
    location = payload.get("location", "").strip()
    if not name:
        conn.close()
        return jsonify({"error": "Office name is required."}), 400
    try:
        cursor.execute("INSERT INTO offices (name, location) VALUES (?, ?)", (name, location))
        conn.commit()
        return jsonify({"success": True})
    except sqlite3.IntegrityError:
        return jsonify({"error": "Office already exists."}), 400
    finally:
        conn.close()

@app.route("/api/offices/<int:office_id>", methods=["DELETE"])
def delete_office(office_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM offices WHERE id = ?", (office_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/designations", methods=["GET", "POST"])
def manage_designations():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM designations ORDER BY title")
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    title = payload.get("title", "").strip()
    department_id = payload.get("department_id")
    if not title:
        conn.close()
        return jsonify({"error": "Designation title is required."}), 400
    cursor.execute("INSERT INTO designations (title, department_id) VALUES (?, ?)", (title, department_id))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/designations/<int:designation_id>", methods=["DELETE"])
def delete_designation(designation_id):
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("DELETE FROM designations WHERE id = ?", (designation_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/shifts", methods=["GET", "POST"])
def manage_shifts():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM shifts ORDER BY name")
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    name = payload.get("name", "").strip()
    start_time = payload.get("start_time", "").strip()
    end_time = payload.get("end_time", "").strip()
    description = payload.get("description", "").strip()
    if not name:
        conn.close()
        return jsonify({"error": "Shift name is required."}), 400
    cursor.execute(
        "INSERT INTO shifts (name, start_time, end_time, description) VALUES (?, ?, ?, ?)",
        (name, start_time, end_time, description)
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/shifts/<int:shift_id>", methods=["GET", "PUT", "DELETE"])
def manage_shift(shift_id):
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM shifts WHERE id = ?", (shift_id,))
        row = cursor.fetchone()
        conn.close()
        return jsonify(dict(row) if row else {})

    if request.method == "PUT":
        payload = request.json or {}
        name = payload.get("name", "").strip()
        start_time = payload.get("start_time", "").strip()
        end_time = payload.get("end_time", "").strip()
        description = payload.get("description", "").strip()
        if not name:
            conn.close()
            return jsonify({"error": "Shift name is required."}), 400
        cursor.execute(
            "UPDATE shifts SET name = ?, start_time = ?, end_time = ?, description = ? WHERE id = ?",
            (name, start_time, end_time, description, shift_id)
        )
        conn.commit()
        conn.close()
        return jsonify({"success": True})

    cursor.execute("DELETE FROM shifts WHERE id = ?", (shift_id,))
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/employees", methods=["GET", "POST"])
def manage_employees():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        data = get_employee_list()
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    required_fields = ["employee_code", "first_name", "last_name", "department_id", "designation_id", "office_id", "joining_date", "ctc"]
    for field in required_fields:
        if not payload.get(field):
            conn.close()
            return jsonify({"error": f"{field.replace('_', ' ').title()} is required."}), 400

    employee_data = (
        payload.get("employee_code", "").strip() or None,
        payload.get("first_name").strip(),
        payload.get("last_name").strip(),
        payload.get("email", "").strip(),
        payload.get("phone", "").strip(),
        payload.get("dob", "").strip() or None,
        payload.get("gender", "").strip(),
        payload.get("department_id"),
        payload.get("designation_id"),
        payload.get("office_id"),
        payload.get("manager_id"),
        payload.get("joining_date", "").strip() or None,
        float(payload.get("ctc") or 0),
        payload.get("status", "Active").strip(),
        payload.get("address", "").strip(),
        int(bool(payload.get("pf_active", 1))),
        int(bool(payload.get("esic_active", 1))),
        int(bool(payload.get("professional_tax_active", 1))),
        int(bool(payload.get("lwf_active", 1)))
    )

    try:
        cursor.execute(
            "INSERT INTO employees (employee_code, first_name, last_name, email, phone, dob, gender, department_id, designation_id, office_id, manager_id, joining_date, ctc, status, address, pf_active, esic_active, professional_tax_active, lwf_active) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            employee_data
        )
        employee_id = cursor.lastrowid
        # Insert default leave balance
        current_year = datetime.now().year
        cursor.execute(
            "INSERT INTO leave_balances (employee_id, earned_leave, casual_leave, loss_of_pay, year) VALUES (?, 12, 6, 0, ?)",
            (employee_id, current_year)
        )
        conn.commit()
        return jsonify({"success": True})
    except sqlite3.IntegrityError:
        return jsonify({"error": "Employee code must be unique."}), 400
    finally:
        conn.close()

@app.route("/api/employees/<int:employee_id>", methods=["GET", "PUT", "DELETE"])
def employee_detail(employee_id):
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM employees WHERE id = ?", (employee_id,))
        row = cursor.fetchone()
        conn.close()
        return jsonify(dict(row) if row else {})

    if request.method == "DELETE":
        cursor.execute("DELETE FROM employees WHERE id = ?", (employee_id,))
        conn.commit()
        conn.close()
        return jsonify({"success": True})

    payload = request.json or {}
    employee_data = (
        payload.get("employee_code", "").strip(),
        payload.get("first_name", "").strip(),
        payload.get("last_name", "").strip(),
        payload.get("email", "").strip(),
        payload.get("phone", "").strip(),
        payload.get("dob", "").strip(),
        payload.get("gender", "").strip(),
        payload.get("department_id"),
        payload.get("designation_id"),
        payload.get("office_id"),
        payload.get("manager_id"),
        payload.get("joining_date", "").strip(),
        float(payload.get("ctc", 0)),
        payload.get("status", "Active").strip(),
        payload.get("address", "").strip(),
        int(bool(payload.get("pf_active", 1))),
        int(bool(payload.get("esic_active", 1))),
        int(bool(payload.get("professional_tax_active", 1))),
        int(bool(payload.get("lwf_active", 1))),
        employee_id
    )
    cursor.execute(
        "UPDATE employees SET employee_code = ?, first_name = ?, last_name = ?, email = ?, phone = ?, dob = ?, gender = ?, department_id = ?, designation_id = ?, office_id = ?, manager_id = ?, joining_date = ?, ctc = ?, status = ?, address = ?, pf_active = ?, esic_active = ?, professional_tax_active = ?, lwf_active = ? WHERE id = ?",
        employee_data
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/employees/<int:employee_id>/statutory", methods=["PUT"])
def update_statutory(employee_id):
    conn = get_connection()
    cursor = conn.cursor()
    payload = request.json or {}
    updates = {}
    if "pf_active" in payload:
        updates["pf_active"] = int(bool(payload["pf_active"]))
    if "esic_active" in payload:
        updates["esic_active"] = int(bool(payload["esic_active"]))
    if "pt_active" in payload:
        updates["pt_active"] = int(bool(payload["pt_active"]))
    if "lwf_active" in payload:
        updates["lwf_active"] = int(bool(payload["lwf_active"]))
    if not updates:
        conn.close()
        return jsonify({"error": "No valid fields to update."}), 400
    set_clause = ", ".join(f"{k} = ?" for k in updates)
    values = list(updates.values()) + [employee_id]
    cursor.execute(f"UPDATE employees SET {set_clause} WHERE id = ?", values)
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/download-template")
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    headers = ["employee_code", "first_name", "last_name", "email", "phone", "dob", "gender", "department_name", "designation_title", "office_name", "manager_name", "joining_date", "ctc", "status", "address"]
    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)
    sample_data = ["EMP001", "John", "Doe", "john@example.com", "1234567890", "1990-01-01", "Male", "Engineering", "Developer", "Head Office", "", "2023-01-01", "50000", "Active", "123 Main St"]
    for col_num, value in enumerate(sample_data, 1):
        ws.cell(row=2, column=col_num, value=value)
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="employee_import_template.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@app.route("/api/employees/bulk-import", methods=["POST"])
def bulk_import_employees():
    if 'file' not in request.files:
        return jsonify({"error": "No file provided"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({"error": "Invalid file type. Please upload an Excel file."}), 400
    
    try:
        wb = load_workbook(file)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return jsonify({"error": "File is empty"}), 400
        headers = [str(h).strip().lower() for h in rows[0]]
        expected_headers = ["employee_code", "first_name", "last_name", "email", "phone", "dob", "gender", "department_name", "designation_title", "office_name", "manager_name", "joining_date", "ctc", "status", "address"]
        if headers != expected_headers:
            return jsonify({"error": "Invalid headers. Please use the template."}), 400
        
        conn = get_connection()
        cursor = conn.cursor()
        imported = 0
        updated = 0
        for row in rows[1:]:
            # Safe access to avoid index errors on short rows
            def g(i, default=None):
                try:
                    return row[i]
                except (IndexError, TypeError):
                    return default

            def parse_str(val):
                s = str(val or '').strip()
                return s if s else None

            emp_code_raw = g(0)
            emp_code = parse_str(emp_code_raw)
            if not emp_code:
                continue

            first = parse_str(g(1))
            last = parse_str(g(2))
            email = parse_str(g(3))
            phone = parse_str(g(4))
            dob = normalize_date_value(g(5))
            gender = parse_str(g(6))
            dept_name = parse_str(g(7))
            des_title = parse_str(g(8))
            off_name = parse_str(g(9))
            mgr_raw = parse_str(g(10))
            joining_date = normalize_date_value(g(11))
            ctc_raw = g(12)
            ctc_val = None
            if ctc_raw is not None:
                try:
                    ctc_str = str(ctc_raw).replace(',', '').replace('$', '').strip()
                    ctc_val = float(ctc_str) if ctc_str else None
                except Exception:
                    ctc_val = None
            status = parse_str(g(13))
            address = parse_str(g(14))

            # Check if employee exists by code
            cursor.execute("SELECT id FROM employees WHERE employee_code = ?", (emp_code,))
            existing = cursor.fetchone()

            if existing:
                # UPDATE only non-blank fields from excel; leave DB values if blank in excel
                set_clauses = []
                params = []

                if first is not None:
                    set_clauses.append("first_name=?")
                    params.append(first)
                if last is not None:
                    set_clauses.append("last_name=?")
                    params.append(last)
                if email is not None:
                    set_clauses.append("email=?")
                    params.append(email)
                if phone is not None:
                    set_clauses.append("phone=?")
                    params.append(phone)
                if dob is not None:
                    set_clauses.append("dob=?")
                    params.append(dob)
                if gender is not None:
                    set_clauses.append("gender=?")
                    params.append(gender)

                if dept_name is not None:
                    department_id = None
                    if dept_name:
                        cursor.execute("SELECT id FROM departments WHERE name = ?", (dept_name,))
                        dep = cursor.fetchone()
                        department_id = dep[0] if dep else None
                    set_clauses.append("department_id=?")
                    params.append(department_id)

                if des_title is not None:
                    designation_id = None
                    if des_title:
                        cursor.execute("SELECT id FROM designations WHERE title = ?", (des_title,))
                        des = cursor.fetchone()
                        designation_id = des[0] if des else None
                    set_clauses.append("designation_id=?")
                    params.append(designation_id)

                if off_name is not None:
                    office_id = None
                    if off_name:
                        cursor.execute("SELECT id FROM offices WHERE name = ?", (off_name,))
                        off = cursor.fetchone()
                        office_id = off[0] if off else None
                    set_clauses.append("office_id=?")
                    params.append(office_id)

                if mgr_raw is not None:
                    manager_id = None
                    if mgr_raw:
                        names = mgr_raw.split()
                        if len(names) >= 2:
                            cursor.execute("SELECT id FROM employees WHERE first_name = ? AND last_name = ?", (names[0], names[1]))
                            mgr = cursor.fetchone()
                            manager_id = mgr[0] if mgr else None
                    set_clauses.append("manager_id=?")
                    params.append(manager_id)

                if joining_date is not None:
                    set_clauses.append("joining_date=?")
                    params.append(joining_date)
                if ctc_val is not None:
                    set_clauses.append("ctc=?")
                    params.append(ctc_val)
                if status is not None:
                    set_clauses.append("status=?")
                    params.append(status)
                if address is not None:
                    set_clauses.append("address=?")
                    params.append(address)

                if set_clauses:
                    params.append(emp_code)
                    cursor.execute(
                        f"UPDATE employees SET {', '.join(set_clauses)} WHERE employee_code = ?",
                        params
                    )
                    updated += 1
            else:
                # INSERT new: require first + last (as before)
                if first is None or last is None:
                    continue

                department_id = None
                if dept_name:
                    cursor.execute("SELECT id FROM departments WHERE name = ?", (dept_name,))
                    dep = cursor.fetchone()
                    department_id = dep[0] if dep else None

                designation_id = None
                if des_title:
                    cursor.execute("SELECT id FROM designations WHERE title = ?", (des_title,))
                    des = cursor.fetchone()
                    designation_id = des[0] if des else None

                office_id = None
                if off_name:
                    cursor.execute("SELECT id FROM offices WHERE name = ?", (off_name,))
                    off = cursor.fetchone()
                    office_id = off[0] if off else None

                manager_id = None
                if mgr_raw:
                    names = mgr_raw.split()
                    if len(names) >= 2:
                        cursor.execute("SELECT id FROM employees WHERE first_name = ? AND last_name = ?", (names[0], names[1]))
                        mgr = cursor.fetchone()
                        manager_id = mgr[0] if mgr else None

                gender = gender or "Male"
                status = status or "Active"
                ctc_val = ctc_val if ctc_val is not None else 0.0

                employee_data = (
                    emp_code,
                    first,
                    last,
                    email,
                    phone,
                    dob,
                    gender,
                    department_id,
                    designation_id,
                    office_id,
                    manager_id,
                    joining_date,
                    ctc_val,
                    status,
                    address,
                    1, 1, 1, 1
                )
                try:
                    cursor.execute(
                        "INSERT INTO employees (employee_code, first_name, last_name, email, phone, dob, gender, department_id, designation_id, office_id, manager_id, joining_date, ctc, status, address, pf_active, esic_active, professional_tax_active, lwf_active) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                        employee_data
                    )
                    employee_id = cursor.lastrowid
                    current_year = datetime.now().year
                    cursor.execute(
                        "INSERT OR IGNORE INTO leave_balances (employee_id, earned_leave, casual_leave, loss_of_pay, year) VALUES (?, 12, 6, 0, ?)",
                        (employee_id, current_year)
                    )
                    imported += 1
                except sqlite3.IntegrityError:
                    continue
        conn.commit()
        conn.close()
        return jsonify({"imported": imported, "updated": updated})
    except Exception as e:
        return jsonify({"error": str(e)}), 500


@app.route("/api/employee-joining-forms", methods=["POST"])
def save_employee_joining_form():
    conn = get_connection()
    cursor = conn.cursor()

    # Generate employee code first
    employee_code = get_next_employee_code()

    # Text fields from form
    full_name = request.form.get("joining_full_name", "")
    contact_no = request.form.get("joining_contact_no", "")
    email_id = request.form.get("joining_email_id", "")
    designation = request.form.get("joining_designation", "")
    joining_date = request.form.get("joining_joining_date", "")
    permanent_address = request.form.get("joining_permanent_address", "")
    date_of_birth = request.form.get("joining_date_of_birth", "")
    gender = request.form.get("joining_gender", "")
    marital_status = request.form.get("joining_marital_status", "")
    pan_no = request.form.get("joining_pan_no", "")
    aadhar_no = request.form.get("joining_aadhar_no", "")
    bank_name = request.form.get("joining_bank_name", "")
    bank_account_number = request.form.get("joining_bank_account_number", "")
    ifsc_code = request.form.get("joining_ifsc_code", "")
    branch_name = request.form.get("joining_branch_name", "")
    highest_qualification = request.form.get("joining_highest_qualification", "")
    university_board = request.form.get("joining_university_board", "")
    year_of_passing = request.form.get("joining_year_of_passing", "")
    percentage_grade = request.form.get("joining_percentage_grade", "")
    emergency_person_name = request.form.get("joining_emergency_person_name", "")
    emergency_person_mobile = request.form.get("joining_emergency_person_mobile", "")
    emergency_relation = request.form.get("joining_emergency_relation", "")
    emergency_person_address = request.form.get("joining_emergency_person_address", "")
    prev_company_name = request.form.get("joining_prev_company_name", "")
    prev_designation = request.form.get("joining_prev_designation", "")
    prev_duration = request.form.get("joining_prev_duration", "")
    prev_last_salary = request.form.get("joining_prev_last_salary", "")

    # Handle file uploads (local + optional Google Drive)
    def save_uploaded_file(file_key):
        if file_key in request.files:
            file = request.files[file_key]
            if file and file.filename:
                filename = secure_filename(file.filename)
                filename = f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{filename}"
                filepath = os.path.join(UPLOAD_FOLDER, filename)
                file.save(filepath)
                
                # Optional Google Drive upload (uses your existing google_drive_manager setup)
                try:
                    from google_drive_manager import GoogleDriveManager
                    gdm = GoogleDriveManager()
                    # Using the existing "Offer Letters" folder for now; documents will be private to your account
                    folder_id = gdm.create_offer_letters_folder()
                    with open(filepath, 'rb') as f:
                        content = f.read()
                    drive_res = gdm.upload_file(filename, content, folder_id=folder_id)
                    print(f"✅ Also uploaded to Google Drive: {drive_res.get('view_link')}")
                except Exception as drive_err:
                    print(f"⚠️ Google Drive upload skipped for {file_key}: {drive_err}")
                
                return filename
        return ""
    photo_filename = save_uploaded_file("joining_photo")
    pan_card_filename = save_uploaded_file("joining_pan_card")
    aadhar_card_filename = save_uploaded_file("joining_aadhar_card")
    cheque_passbook_filename = save_uploaded_file("joining_cheque_passbook")
    highest_education_cert_filename = save_uploaded_file("joining_highest_education_cert")
    last_3_month_salary_slip_filename = save_uploaded_file("joining_last_3_month_salary_slip")
    prev_employment_docs_filename = save_uploaded_file("joining_prev_employment_docs")

    cursor.execute(
        "INSERT INTO employee_joining_forms ("
        "full_name, contact_no, email_id, designation, joining_date, permanent_address, date_of_birth, gender, marital_status, "
        "pan_no, aadhar_no, bank_name, bank_account_number, ifsc_code, branch_name, "
        "highest_qualification, university_board, year_of_passing, percentage_grade, "
        "emergency_person_name, emergency_person_mobile, emergency_relation, emergency_person_address, "
        "prev_company_name, prev_designation, prev_duration, prev_last_salary, "
        "photo_filename, pan_card_filename, aadhar_card_filename, cheque_passbook_filename, "
        "highest_education_cert_filename, last_3_month_salary_slip_filename, prev_employment_docs_filename, "
        "employee_code, status, submitted_on"
        ") VALUES ("
        "?, ?, ?, ?, ?, ?, ?, ?, ?, "
        "?, ?, ?, ?, ?, ?, "
        "?, ?, ?, ?, "
        "?, ?, ?, ?, "
        "?, ?, ?, ?, "
        "?, ?, ?, ?, "
        "?, ?, ?, "
        "?, ?, ?"
        ")",
        (
            full_name, contact_no, email_id, designation, joining_date, permanent_address, date_of_birth, gender, marital_status,
            pan_no, aadhar_no, bank_name, bank_account_number, ifsc_code, branch_name,
            highest_qualification, university_board, year_of_passing, percentage_grade,
            emergency_person_name, emergency_person_mobile, emergency_relation, emergency_person_address,
            prev_company_name, prev_designation, prev_duration, prev_last_salary,
            photo_filename, pan_card_filename, aadhar_card_filename, cheque_passbook_filename,
            highest_education_cert_filename, last_3_month_salary_slip_filename, prev_employment_docs_filename,
            employee_code, 'pending',
            datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        )
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True, "employee_code": employee_code})


@app.route("/api/reject-joining", methods=["POST"])
def reject_joining():
    payload = request.json or {}
    joining_id = payload.get("id")
    if not joining_id:
        return jsonify({"error": "Joining ID required"}), 400
    
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "UPDATE employee_joining_forms SET status = 'rejected' WHERE id = ?",
        (joining_id,)
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@app.route("/api/pending-joinings")
def get_pending_joinings():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("""
        SELECT * FROM employee_joining_forms 
        WHERE status = 'pending' 
        ORDER BY submitted_on DESC
    """)
    rows = cursor.fetchall()
    conn.close()
    return jsonify([dict(row) for row in rows])


@app.route("/api/approve-joining", methods=["POST"])
def approve_joining():
    payload = request.json or {}
    joining_id = payload.get("id")
    if not joining_id:
        return jsonify({"error": "Joining ID required"}), 400
    
    conn = get_connection()
    cursor = conn.cursor()
    
    # Get the pending record
    cursor.execute("SELECT * FROM employee_joining_forms WHERE id = ?", (joining_id,))
    joining = cursor.fetchone()
    if not joining or joining['status'] != 'pending':
        conn.close()
        return jsonify({"error": "Invalid or already processed joining"}), 400
    
    # Split full name for employees table
    full_name = joining['full_name'] or ''
    parts = full_name.split(' ', 1)
    first_name = parts[0] if parts else ''
    last_name = parts[1] if len(parts) > 1 else ''
    
    # Map fields (some will be NULL/default since joining form has limited data)
    employee_data = (
        joining['employee_code'],
        first_name,
        last_name,
        joining['email_id'],
        joining['contact_no'],
        joining['date_of_birth'],
        joining['gender'],
        None,  # department_id - admin can update later
        None,  # designation_id - use text for now or map
        None,  # office_id
        None,  # manager_id
        joining['joining_date'],
        0,     # ctc default
        'Active',
        joining['permanent_address'],
        1, 1, 1, 1
    )
    
    try:
        cursor.execute(
            "INSERT INTO employees (employee_code, first_name, last_name, email, phone, dob, gender, department_id, designation_id, office_id, manager_id, joining_date, ctc, status, address, pf_active, esic_active, professional_tax_active, lwf_active) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            employee_data
        )
        
        # Update joining status
        cursor.execute(
            "UPDATE employee_joining_forms SET status = 'approved' WHERE id = ?",
            (joining_id,)
        )
        
        # Also create default leave balance for the new employee
        employee_id = cursor.lastrowid
        current_year = datetime.now().year
        cursor.execute(
            "INSERT OR IGNORE INTO leave_balances (employee_id, earned_leave, casual_leave, loss_of_pay, year) VALUES (?, 12, 6, 0, ?)",
            (employee_id, current_year)
        )
        
        conn.commit()
        conn.close()
        return jsonify({"success": True})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({"error": str(e)}), 500


@app.route("/download-attendance-template")
def download_attendance_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    headers = ["employee_code", "date", "status", "shift_name", "check_in", "check_out"]
    for col_num, header in enumerate(headers, 1):
        ws.cell(row=1, column=col_num, value=header)
    sample_data = ["EMP001", "2026-04-14", "P", "Day Shift", "09:00", "18:00"]
    for col_num, value in enumerate(sample_data, 1):
        ws.cell(row=2, column=col_num, value=value)
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name="attendance_import_template.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

def normalize_date_value(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().strftime('%Y-%m-%d')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')

    text = str(value).strip()
    if not text:
        return None

    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%d.%m.%Y', '%d %b %Y', '%d %B %Y'):
        try:
            return datetime.strptime(text, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue

    return text

@app.route("/api/attendance/bulk-import", methods=["POST"])
def bulk_import_attendance():
    if 'file' not in request.files:
        return jsonify({"error": "No file provided"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
    if not file.filename.endswith(('.xlsx', '.xls')):
        return jsonify({"error": "Invalid file type. Please upload an Excel file."}), 400
    try:
        wb = load_workbook(file)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return jsonify({"error": "File is empty"}), 400
        headers = [str(h).strip().lower() for h in rows[0]]
        expected_headers = ["employee_code", "date", "status", "shift_name", "check_in", "check_out"]
        if headers != expected_headers:
            return jsonify({"error": "Invalid headers. Please use the attendance template."}), 400
        conn = get_connection()
        cursor = conn.cursor()
        imported = 0
        for row in rows[1:]:
            if not row[0] or not row[1] or not row[2]:
                continue
            cursor.execute("SELECT id FROM employees WHERE employee_code = ?", (str(row[0]).strip(),))
            employee = cursor.fetchone()
            if not employee:
                continue
            employee_id = employee[0]
            shift_id = None
            if row[3]:
                cursor.execute("SELECT id FROM shifts WHERE name = ?", (str(row[3]).strip(),))
                shift = cursor.fetchone()
                shift_id = shift[0] if shift else None
            date_value = normalize_date_value(row[1])
            if not date_value:
                continue
            status_value = str(row[2]).strip()
            check_in = str(row[4]).strip() if row[4] else None
            check_out = str(row[5]).strip() if row[5] else None
            cursor.execute("SELECT id FROM attendance WHERE employee_id = ? AND date = ?", (employee_id, date_value))
            existing = cursor.fetchone()
            if existing:
                cursor.execute(
                    "UPDATE attendance SET status = ?, shift_id = ?, check_in = ?, check_out = ? WHERE id = ?",
                    (status_value, shift_id, check_in, check_out, existing[0])
                )
            else:
                cursor.execute(
                    "INSERT INTO attendance (employee_id, date, status, shift_id, check_in, check_out) VALUES (?, ?, ?, ?, ?, ?)",
                    (employee_id, date_value, status_value, shift_id, check_in, check_out)
                )
            imported += 1
        conn.commit()
        conn.close()
        return jsonify({"imported": imported})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@app.route("/api/leave-balances")
def get_leave_balances():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT lb.*, e.first_name || ' ' || e.last_name AS employee_name FROM leave_balances lb "
        "LEFT JOIN employees e ON lb.employee_id = e.id "
        "WHERE EXISTS (SELECT 1 FROM leave_requests lr WHERE lr.employee_id = lb.employee_id) "
        "ORDER BY e.first_name, e.last_name"
    )
    data = dict_rows(cursor.fetchall())
    conn.close()
    return jsonify(data)

@app.route("/api/statutory")
def get_statutory():
    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id, first_name || ' ' || last_name AS name, pf_active, esic_active, professional_tax_active AS pt_active, lwf_active FROM employees "
        "ORDER BY first_name, last_name"
    )
    data = dict_rows(cursor.fetchall())
    conn.close()
    return jsonify(data)

@app.route("/api/attendance", methods=["GET", "POST"])
def manage_attendance():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        employee_id = request.args.get('employee_id')
        start_date = request.args.get('start_date')
        end_date = request.args.get('end_date')
        query = """
            SELECT a.*, e.first_name || ' ' || e.last_name AS employee_name, s.name AS shift_name FROM attendance a 
            LEFT JOIN employees e ON a.employee_id = e.id 
            LEFT JOIN shifts s ON a.shift_id = s.id 
        """
        conditions = []
        params = []
        if employee_id:
            conditions.append("a.employee_id = ?")
            params.append(employee_id)
        fetch_start = start_date
        if start_date:
            try:
                start_dt = datetime.strptime(start_date, '%Y-%m-%d').date()
                monday = start_dt - timedelta(days=start_dt.weekday())
                fetch_start = monday.strftime('%Y-%m-%d')
            except Exception:
                pass
        if fetch_start:
            conditions.append("a.date >= ?")
            params.append(fetch_start)
        if end_date:
            conditions.append("a.date <= ?")
            params.append(end_date)
        if conditions:
            query += " WHERE " + " AND ".join(conditions)
        query += " ORDER BY a.date DESC"
        cursor.execute(query, params)
        data = dict_rows(cursor.fetchall())
        data = apply_weekend_absence_rule(data, start_date=start_date, end_date=end_date)
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    employee_id = payload.get('employee_id')
    date = payload.get('date')
    status = payload.get('status')
    shift_id = payload.get('shift_id')
    check_in = payload.get('check_in')
    check_out = payload.get('check_out')
    remarks = payload.get('remarks', '')

    required_fields = ["employee_id", "date", "status"]
    for field in required_fields:
        if not payload.get(field):
            conn.close()
            return jsonify({"error": f"{field.replace('_', ' ').title()} is required."}), 400

    cursor.execute(
        "INSERT INTO attendance (employee_id, date, status, shift_id, check_in, check_out, remarks) VALUES (?, ?, ?, ?, ?, ?, ?)",
        (employee_id, date, status, shift_id, check_in, check_out, remarks)
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/attendance/<int:attendance_id>", methods=["PUT"])
def update_attendance(attendance_id):
    payload = request.json or {}
    status = payload.get('status')
    shift_id = payload.get('shift_id')
    check_in = payload.get('check_in')
    check_out = payload.get('check_out')
    remarks = payload.get('remarks')

    conn = get_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT * FROM attendance WHERE id = ?", (attendance_id,))
    existing = cursor.fetchone()
    if not existing:
        conn.close()
        return jsonify({"error": "Attendance record not found."}), 404

    cursor.execute(
        "UPDATE attendance SET status = ?, shift_id = ?, check_in = ?, check_out = ?, remarks = ? WHERE id = ?",
        (
            status if status is not None else existing['status'],
            shift_id if shift_id is not None else existing['shift_id'],
            check_in if check_in is not None else existing['check_in'],
            check_out if check_out is not None else existing['check_out'],
            remarks if remarks is not None else existing.get('remarks'),
            attendance_id
        )
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})


@app.route("/api/attendance/bulk", methods=["POST"])
def bulk_mark_attendance():
    payload = request.json or {}
    date = payload.get('date')
    records = payload.get('records') or []
    if not date:
        return jsonify({"error": "date is required."}), 400
    if not isinstance(records, list) or len(records) == 0:
        return jsonify({"error": "No attendance records provided."}), 400

    conn = get_connection()
    cursor = conn.cursor()
    saved = 0
    for rec in records:
        emp_id = rec.get('employee_id')
        status = rec.get('status')
        remarks = rec.get('remarks', '')
        if not emp_id or not status:
            continue
        cursor.execute("SELECT id FROM attendance WHERE employee_id = ? AND date = ?", (emp_id, date))
        existing = cursor.fetchone()
        if existing:
            cursor.execute(
                "UPDATE attendance SET status = ?, remarks = ? WHERE id = ?",
                (status, remarks, existing[0])
            )
        else:
            cursor.execute(
                "INSERT INTO attendance (employee_id, date, status, remarks) VALUES (?, ?, ?, ?)",
                (emp_id, date, status, remarks)
            )
        saved += 1
    conn.commit()
    conn.close()
    return jsonify({"success": True, "saved": saved})


@app.route("/api/attendance/summary")
def attendance_summary():
    conn = get_connection()
    cursor = conn.cursor()
    employee_id = request.args.get('employee_id')
    start_date = request.args.get('start_date')
    end_date = request.args.get('end_date')

    today = datetime.now().date()
    if not end_date:
        end_date = today.strftime('%Y-%m-%d')
    if not start_date:
        start_date = (today - timedelta(days=6)).strftime('%Y-%m-%d')

    start_dt = datetime.strptime(start_date, '%Y-%m-%d').date()
    end_dt = datetime.strptime(end_date, '%Y-%m-%d').date()
    if start_dt > end_dt:
        start_dt, end_dt = end_dt, start_dt

    date_labels = []
    current = start_dt
    while current <= end_dt:
        date_labels.append(current.strftime('%d %a'))
        current += timedelta(days=1)

    if employee_id:
        cursor.execute(
            "SELECT id, first_name || ' ' || last_name AS employee_name FROM employees WHERE id = ? ORDER BY first_name, last_name",
            (employee_id,)
        )
    else:
        cursor.execute(
            "SELECT id, first_name || ' ' || last_name AS employee_name FROM employees ORDER BY first_name, last_name"
        )
    employees = dict_rows(cursor.fetchall())

    fetch_start = start_date
    try:
        start_dt_obj = datetime.strptime(start_date, '%Y-%m-%d').date()
        monday = start_dt_obj - timedelta(days=start_dt_obj.weekday())
        fetch_start = monday.strftime('%Y-%m-%d')
    except Exception:
        pass

    query = "SELECT a.id, a.employee_id, a.date, a.status, s.name AS shift_name FROM attendance a LEFT JOIN shifts s ON a.shift_id = s.id WHERE a.date BETWEEN ? AND ?"
    params = [fetch_start, end_date]
    if employee_id:
        query += " AND a.employee_id = ?"
        params.append(employee_id)
    cursor.execute(query, params)
    attendance_rows = dict_rows(cursor.fetchall())
    attendance_rows = apply_weekend_absence_rule(attendance_rows, start_date=start_date, end_date=end_date)

    attendance_map = {}
    for row in attendance_rows:
        if row['employee_id'] not in attendance_map:
            attendance_map[row['employee_id']] = {}
        attendance_map[row['employee_id']][row['date']] = row

    rows = []
    for emp in employees:
        row = {
            'id': emp['id'],
            'employee_name': emp['employee_name'],
            'present': 0,
            'absent': 0,
            'week_off': 0,
            'holiday': 0,
            'half_day': 0,
            'total': 0,
            'daily': []
        }
        daily = []
        for i in range((end_dt - start_dt).days + 1):
            date = start_dt + timedelta(days=i)
            date_key = date.strftime('%Y-%m-%d')
            record = attendance_map.get(emp['id'], {}).get(date_key)
            if record:
                status = record['status']
                row['total'] += 1
                if status == 'P':
                    row['present'] += 1
                elif status == 'A':
                    row['absent'] += 1
                elif status == 'WO':
                    row['week_off'] += 1
                elif status == 'H':
                    row['holiday'] += 1
                elif status in ('AH', 'HF'):
                    row['half_day'] += 1
                daily.append({
                    'date': date_key,
                    'status': status,
                    'attendance_id': record['id']
                })
            else:
                daily.append({
                    'date': date_key,
                    'status': '',
                    'attendance_id': ''
                })
        row['daily'] = daily
        rows.append(row)

    conn.close()
    return jsonify({'headers': date_labels, 'rows': rows})

@app.route("/api/leaves", methods=["GET", "POST"])
def manage_leaves():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute(
            "SELECT l.*, e.first_name || ' ' || e.last_name AS employee_name FROM leave_requests l "
            "LEFT JOIN employees e ON l.employee_id = e.id "
            "ORDER BY l.applied_on DESC"
        )
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    required_fields = ["employee_id", "leave_type", "start_date", "end_date"]
    for field in required_fields:
        if not payload.get(field):
            conn.close()
            return jsonify({"error": f"{field.replace('_', ' ').title()} is required."}), 400

    cursor.execute(
        "INSERT INTO leave_requests (employee_id, leave_type, start_date, end_date, status, reason, applied_on) VALUES (?, ?, ?, ?, ?, ?, ?)",
        (
            payload.get("employee_id"),
            payload.get("leave_type"),
            payload.get("start_date"),
            payload.get("end_date"),
            payload.get("status", "Pending"),
            payload.get("reason", ""),
            datetime.now().strftime("%Y-%m-%d")
        )
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/leaves/<int:leave_id>", methods=["GET", "PUT", "DELETE"])
def manage_leave_request(leave_id):
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute("SELECT * FROM leave_requests WHERE id = ?", (leave_id,))
        row = cursor.fetchone()
        conn.close()
        return jsonify(dict(row) if row else {})

    if request.method == "DELETE":
        cursor.execute("DELETE FROM leave_requests WHERE id = ?", (leave_id,))
        conn.commit()
        conn.close()
        return jsonify({"success": True})

    cursor.execute("SELECT * FROM leave_requests WHERE id = ?", (leave_id,))
    existing = cursor.fetchone()
    if not existing:
        conn.close()
        return jsonify({"error": "Leave request not found."}), 404

    payload = request.json or {}
    employee_id = payload.get("employee_id", existing["employee_id"])
    leave_type = payload.get("leave_type", existing["leave_type"]).strip()
    start_date = payload.get("start_date", existing["start_date"]).strip()
    end_date = payload.get("end_date", existing["end_date"]).strip()
    status = payload.get("status", existing["status"]).strip()
    reason = payload.get("reason", existing["reason"]).strip()

    cursor.execute(
        "UPDATE leave_requests SET employee_id = ?, leave_type = ?, start_date = ?, end_date = ?, status = ?, reason = ? WHERE id = ?",
        (employee_id, leave_type, start_date, end_date, status, reason, leave_id)
    )
    conn.commit()
    conn.close()
    return jsonify({"success": True})

@app.route("/api/credit-earned-leaves/<int:year>/<int:month>", methods=["POST"])
def credit_earned_leaves(year, month):
    conn = get_connection()
    cursor = conn.cursor()
    
    # Get all employees
    cursor.execute("SELECT id FROM employees")
    employees = cursor.fetchall()
    
    credited_count = 0
    for emp in employees:
        emp_id = emp[0]
        
        # Check if already credited
        cursor.execute(
            "SELECT id FROM earned_leave_credits WHERE employee_id = ? AND year = ? AND month = ?",
            (emp_id, year, month)
        )
        if cursor.fetchone():
            continue  # Already credited
        
        # Calculate present days
        start_date = date(year, month, 1)
        if month == 12:
            end_date = date(year + 1, 1, 1) - timedelta(days=1)
        else:
            end_date = date(year, month + 1, 1) - timedelta(days=1)
        
        present_days = 0.0
        current = start_date
        while current <= end_date:
            cursor.execute(
                "SELECT status FROM attendance WHERE employee_id = ? AND date = ?",
                (emp_id, current.strftime('%Y-%m-%d'))
            )
            row = cursor.fetchone()
            if row:
                status = row[0]
                if status in ('P', 'H', 'WO'):
                    present_days += 1.0
                elif status in ('AH', 'HF'):
                    present_days += 0.5
            current += timedelta(days=1)
        
        if present_days >= 20:
            # Credit 1.25
            cursor.execute(
                "INSERT INTO earned_leave_credits (employee_id, year, month, credited, credited_on) VALUES (?, ?, ?, ?, ?)",
                (emp_id, year, month, 1.25, datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
            )
            # Update leave balance
            cursor.execute(
                "UPDATE leave_balances SET earned_leave = earned_leave + 1.25 WHERE employee_id = ? AND year = ?",
                (emp_id, year)
            )
            credited_count += 1
    
    conn.commit()
    conn.close()
    return jsonify({"success": True, "credited_count": credited_count})

@app.route("/api/payrolls", methods=["GET", "POST"])
def manage_payrolls():
    conn = get_connection()
    cursor = conn.cursor()
    if request.method == "GET":
        cursor.execute(
            "SELECT p.*, e.first_name || ' ' || e.last_name AS employee_name FROM payrolls p "
            "LEFT JOIN employees e ON p.employee_id = e.id "
            "ORDER BY p.year DESC, p.month DESC"
        )
        data = dict_rows(cursor.fetchall())
        conn.close()
        return jsonify(data)

    payload = request.json or {}
    employee_id = payload.get("employee_id")
    month = payload.get("month")
    year = payload.get("year")
    if not employee_id or not month or not year:
        conn.close()
        return jsonify({"error": "Employee, month and year are required."}), 400

    cursor.execute("SELECT * FROM employees WHERE id = ?", (employee_id,))
    employee = cursor.fetchone()
    if not employee:
        conn.close()
        return jsonify({"error": "Employee not found."}), 404

    ctc = float(employee["ctc"] or 0)
    pf_active = bool(employee["pf_active"])
    esic_active = bool(employee["esic_active"])
    professional_tax_active = bool(employee["professional_tax_active"])
    lwf_active = bool(employee["lwf_active"])
    payroll = calculate_pay_components(ctc, pf_active, esic_active, professional_tax_active, lwf_active)
    cursor.execute(
        "SELECT id FROM payrolls WHERE employee_id = ? AND month = ? AND year = ?",
        (employee_id, month, year)
    )
    if cursor.fetchone():
        conn.close()
        return jsonify({"error": "Payroll already generated for this employee and period."}), 400

    cursor.execute(
        "INSERT INTO payrolls (employee_id, month, year, ctc, basic, da, hra, conveyance, special_allowance, gross, pf, esic, professional_tax, lwf, total_deductions, net_pay) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (
            employee_id,
            month,
            year,
            ctc,
            payroll["basic"],
            payroll["da"],
            payroll["hra"],
            payroll["conveyance"],
            payroll["special_allowance"],
            payroll["gross"],
            payroll["pf"],
            payroll["esic"],
            payroll["professional_tax"],
            payroll["lwf"],
            payroll["total_deductions"],
            payroll["net_pay"]
        )
    )
    conn.commit()
    payroll_id = cursor.lastrowid
    cursor.execute("SELECT * FROM payrolls WHERE id = ?", (payroll_id,))
    payroll_record = dict(cursor.fetchone())
    conn.close()
    payroll_record["employee_name"] = f"{employee['first_name']} {employee['last_name']}"
    return jsonify(payroll_record)

@app.route("/generate-offer", methods=["POST"])
def generate_offer():
    payload = request.json or {}
    required = ["offer_date", "prefix", "name", "position", "joining_date", "location", "department", "monthly_ctc"]
    for field in required:
        if not payload.get(field):
            return jsonify({"error": f"{field.replace('_', ' ').title()} is required."}), 400

    ensure_template()
    monthly_ctc = float(payload.get("monthly_ctc"))
    annual_ctc = monthly_ctc * 12
    replacements = {
        "[DATE]": payload.get("offer_date"),
        "[PREFIX]": payload.get("prefix"),
        "[NAME]": payload.get("name"),
        "[POSITION]": payload.get("position"),
        "[JOINING_DATE]": payload.get("joining_date"),
        "[LOCATION]": payload.get("location"),
        "[DEPARTMENT]": payload.get("department"),
        "[CTC_A]": str(int(annual_ctc)),
        "[CTC_A_WORD]": num2words(int(annual_ctc), lang="en_IN").title() + " Only",
        "[CTC_M]": str(int(monthly_ctc)),
        "[CTC_M_WORD]": num2words(int(monthly_ctc), lang="en_IN").title() + " Only"
    }
    bold_keys = ["[NAME]", "[POSITION]", "[JOINING_DATE]", "[LOCATION]", "[DEPARTMENT]", "[CTC_A]", "[CTC_A_WORD]", "[CTC_M]", "[CTC_M_WORD]"]
    temp_path = os.path.join(BASE_DIR, f"offer_letter_{datetime.now().strftime('%Y%m%d%H%M%S')}.docx")
    shutil.copy(TEMPLATE_PATH, temp_path)
    doc = Document(temp_path)
    replace_placeholders(doc, replacements, bold_keys)
    doc.save(temp_path)
    with open(temp_path, 'rb') as f:
        file_bytes = f.read()
    os.remove(temp_path)
    return send_file(
        io.BytesIO(file_bytes),
        mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
        as_attachment=True,
        download_name=f"{payload.get('name')}_Offer_Letter_{payload.get('position')}.docx"
    )

if __name__ == '__main__':
    init_db()
    ensure_template()
    app.run(debug=False, port=5000)
